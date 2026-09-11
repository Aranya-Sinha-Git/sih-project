"""Validate and atomically import operational CSV/XLSX records.

The importer intentionally uses the same AnalyzeInput and analyzed_results
pipeline as the API. Historical similarity is deferred until detail reads so
large imports do not perform a corpus scan once per row.
"""
from __future__ import annotations

import csv
import hashlib
import sys
from pathlib import Path
from typing import Any, Iterable

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.main import AnalyzeInput, _normalized_narrative, analyzed_results, build_incident_record, con, init_db
from app.services.database import get_database

TYPE_ALIASES = {
    "ua": "Unsafe Act",
    "unsafe act": "Unsafe Act",
    "uc": "Unsafe Condition",
    "unsafe condition": "Unsafe Condition",
    "near miss": "Near Miss",
    "incident": "Incident",
}


def _cell(row: dict[str, Any], *names: str) -> Any:
    for name in names:
        if name in row and row[name] not in (None, ""):
            return row[name]
    return ""


def _rows(path: Path) -> Iterable[dict[str, Any]]:
    if path.suffix.casefold() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(values)]
        for values_row in values:
            yield {header: value for header, value in zip(headers, values_row)}
        workbook.close()
        return
    with path.open(newline="", encoding="utf-8-sig") as handle:
        yield from csv.DictReader(handle)


def _payload(row: dict[str, Any]) -> dict[str, Any]:
    raw_type = str(_cell(row, "report_type", "ReportType", "type", "Type") or "Unspecified").strip()
    return {
        "report_id": str(_cell(row, "report_id", "ReportId", "id", "ID") or "").strip() or None,
        "report_date": _cell(row, "report_date", "ReportDate", "date", "Date"),
        "site": str(_cell(row, "site", "Site") or "Unspecified").strip() or "Unspecified",
        "activity": str(_cell(row, "activity", "Activity") or "").strip() or None,
        "report_type": TYPE_ALIASES.get(raw_type.casefold(), raw_type or "Unspecified"),
        "narrative": str(_cell(row, "narrative", "Narrative") or "").strip(),
        "source": str(_cell(row, "source", "Source") or "").strip() or "operational_import",
    }


def main(path: str | Path) -> dict[str, Any]:
    init_db()
    source_path = Path(path)
    if not source_path.exists():
        raise SystemExit(f"File not found: {source_path}")
    raw_rows = list(_rows(source_path))
    reports: list[AnalyzeInput] = []
    errors: list[str] = []
    seen_ids: set[str] = set()
    seen_keys: set[tuple[str, str]] = set()
    for row_number, row in enumerate(raw_rows, start=2):
        try:
            payload = _payload(row)
            if not payload["report_id"]:
                raise ValueError("report_id is required for CLI imports")
            report = AnalyzeInput.model_validate(payload)
            key = (_normalized_narrative(report.narrative), str(report.site or "Unspecified").casefold())
            if report.report_id in seen_ids:
                raise ValueError("duplicate report_id within import")
            if key in seen_keys:
                raise ValueError("duplicate normalized narrative/site within import")
            seen_ids.add(report.report_id)
            seen_keys.add(key)
            reports.append(report)
        except (TypeError, ValueError) as error:
            errors.append(f"Row {row_number}: {error}")
    if errors:
        summary = {"processed": len(raw_rows), "imported": 0, "invalid_or_duplicate": len(errors), "errors": errors}
        print(f"Processed={summary['processed']} imported=0 invalid_or_duplicate={summary['invalid_or_duplicate']} batch=NONE")
        return summary

    database = get_database()
    existing = database.list_incidents()
    existing_ids = {str(item.get("id")) for item in existing}
    existing_keys = {(_normalized_narrative(str(item.get("narrative") or "")), str(item.get("site") or "Unspecified").casefold()) for item in existing}
    batch_id = "IMP-" + hashlib.sha1(source_path.resolve().read_bytes()).hexdigest()[:10].upper()
    analyses = analyzed_results(reports)
    records = []
    skipped = 0
    for report, analysis in zip(reports, analyses):
        key = (_normalized_narrative(report.narrative), str(report.site or "Unspecified").casefold())
        if report.report_id in existing_ids or key in existing_keys:
            skipped += 1
            continue
        analysis["provenance"] = {
            "source_type": "operational_import",
            "source_report_id": report.report_id,
            "import_batch_id": batch_id,
            "source_reference": report.source,
        }
        record, _ = build_incident_record(
            narrative=report.narrative,
            site=report.site or "Unspecified",
            activity=report.activity,
            report_type=report.report_type,
            analysis=analysis,
            report_id=report.report_id,
            report_date=report.report_date,
            source=report.source or "operational_import",
            import_batch_id=batch_id,
        )
        records.append(record)
        existing_ids.add(report.report_id or "")
        existing_keys.add(key)

    connection = con() if database.kind == "sqlite-test" else None
    try:
        if records:
            if database.kind == "sqlite-test":
                database.insert_incidents_batch(records, connection=connection)
                connection.commit()
            else:
                database.insert_incidents_batch(records)
    except Exception:
        if connection is not None:
            connection.rollback()
        raise
    finally:
        if connection is not None:
            connection.close()
    summary = {"processed": len(raw_rows), "imported": len(records), "invalid_or_duplicate": skipped, "batch_id": batch_id, "errors": []}
    print(f"Processed={summary['processed']} imported={summary['imported']} invalid_or_duplicate={summary['invalid_or_duplicate']} batch={batch_id}")
    return summary


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: python scripts/import_incidents.py /path/to/incidents.csv|xlsx")
    main(sys.argv[1])
