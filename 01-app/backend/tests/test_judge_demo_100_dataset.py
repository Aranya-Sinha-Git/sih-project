import csv
import hashlib
import json
import re
from collections import Counter
from datetime import date
from pathlib import Path

DEMO_DIR = Path(__file__).resolve().parents[3] / "04-data" / "demo-datasets" / "judge_demo_100_v0_1"


def test_judge_demo_100_schema_quotas_and_narratives():
    rows = list(csv.DictReader((DEMO_DIR / "judge_demo_100.csv").open(encoding="utf-8")))
    assert len(rows) == 100
    assert list(rows[0]) == ["report_id", "report_date", "site", "activity", "report_type", "narrative", "source"]
    assert [row["report_id"] for row in rows] == [f"DEMO100-{index:04d}" for index in range(1, 101)]
    assert Counter(row["report_type"] for row in rows) == Counter({"Near Miss": 36, "Unsafe Condition": 30, "Unsafe Act": 18, "Incident": 16})
    assert set(row["activity"] for row in rows) == {"Valve Maintenance", "Mechanical Lifting", "Drilling", "Pressure Testing", "Driving", "Hot Work", "Inspection", "Electrical Maintenance"}
    assert len({row["site"] for row in rows}) == 6
    assert all(70 <= len(re.findall(r"\b\w+[\w/-]*\b", row["narrative"])) <= 140 for row in rows)
    assert all(date.fromisoformat(row["report_date"]) <= date.today() for row in rows)
    assert all(row["source"] == "synthetic_judge_demo_100_v0_1" for row in rows)
    assert len({(" ".join(row["narrative"].casefold().split()), row["site"].casefold()) for row in rows}) == 100


def test_judge_demo_100_xlsx_and_release_sidecars():
    from openpyxl import load_workbook

    xlsx = DEMO_DIR / "outputs" / "judge_demo_100_20260912" / "judge_demo_100.xlsx"
    workbook = load_workbook(xlsx, read_only=True, data_only=True)
    sheet = workbook.active
    values = list(sheet.iter_rows(values_only=True))
    assert len(values) == 101 and all(len(row) == 7 for row in values)
    assert sheet["A1"].value == "report_id"
    assert sheet["B2"].value.date().isoformat() <= date.today().isoformat()
    workbook.close()

    profile = json.loads((DEMO_DIR / "judge_demo_100_runtime_profile.json").read_text(encoding="utf-8"))
    manifest = json.loads((DEMO_DIR / "judge_demo_100_manifest.json").read_text(encoding="utf-8"))
    assert profile["status"] == "PASS"
    assert profile["observed"]["risk_distribution"] == {"High": 32, "Medium": 24, "Low": 44}
    assert manifest["status"] == "synthetic_non_validation" and manifest["row_count"] == 100
    assert manifest["duplicate_checks"] == {"duplicate_normalized_narrative_site_pairs": 0, "duplicate_report_ids": 0}
    for filename, details in manifest["files"].items():
        assert hashlib.sha256((DEMO_DIR / filename).read_bytes()).hexdigest() == details["sha256"]
